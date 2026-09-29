"""Excel da lista completa do sistema com IMAGEM (30x30) antes do código + abas de apoio.

Abas: PRODUTOS · SEM FOTO (BUSCAR) · DIVERGÊNCIAS · CATÁLOGO FORA DA LISTA · IMAGENS ENVIADAS · RESUMO · LEIA-ME

Uso: python 07_excel_lista_completa.py <WORK_PLANILHA> <SAIDA_PLANILHA> <ARQUIVO_XLSX> <NOME_PLANILHA_ORIGEM>
"""
import io
import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

WORK, SAIDA, XLSX, ORIGEM = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]), sys.argv[4]
linhas = json.load(open(WORK / "cruzamento.json"))
base = {b["codigo"]: b for b in json.load(open(WORK / "base_produtos.json"))}
fora = json.load(open(WORK / "fora_da_lista.json"))
enviadas = json.load(open(WORK / "sugestoes_imagens_enviadas.json"))

FONTE = "Arial"
AZUL, VERMELHO, CINZA = "10284F", "E81820", "F3F5F9"
PX = 9525
MINI = 30
cab_font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
normal = Font(name=FONTE, size=9)
negrito = Font(name=FONTE, size=9, bold=True)
vermelho = Font(name=FONTE, size=9, color="B00020")
quebra = Alignment(vertical="center", wrap_text=True)
centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
borda = Border(bottom=Side(style="thin", color="D9DEE7"))
COR_STATUS = {"COM FOTO": "D7F0DD", "FOTO DE PRODUTO EQUIVALENTE": "FFF2CC", "FOTO NÃO CONFERE": "F8D7DA",
              "SEM FOTO": "EDEDED"}
INFORMATIVOS = {"Referência vinda do catálogo", "Marca vinda do catálogo"}
PRIORIDADE = ["Marca diferente", "Descrição diferente para o mesmo código", "Código repetido no catálogo",
              "Referência diferente", "Mesma descrição com outro código", "Foto de produto equivalente",
              "Referência vinda do catálogo", "Marca vinda do catálogo"]


def norm(s):
    return unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().upper()


def ordem_natural(t):
    return [(0, float(x.replace(",", "."))) if re.fullmatch(r"\d+(?:[.,]\d+)?", x) else (1, x)
            for x in re.findall(r"\d+(?:[.,]\d+)?|[^\d]+", norm(t))]


_mini = {}


def miniatura(mestre):
    if mestre not in _mini:
        im = Image.open(SAIDA / "_mestres" / mestre).convert("RGBA")
        im.thumbnail((112, 112), Image.LANCZOS)
        tela = Image.new("RGBA", (120, 120), (255, 255, 255, 0))
        tela.alpha_composite(im, ((120 - im.width) // 2, (120 - im.height) // 2))
        buf = io.BytesIO()
        tela.save(buf, format="PNG", optimize=True)
        _mini[mestre] = buf.getvalue()
    return io.BytesIO(_mini[mestre])


def por_imagem(ws, mestre, linha, larg_col):
    img = XLImage(miniatura(mestre))
    col_px = int(larg_col * 7 + 5)
    ox, oy = max(0, (col_px - MINI) // 2) * PX, 3 * PX
    img.anchor = TwoCellAnchor(editAs="twoCell", _from=AnchorMarker(col=0, colOff=ox, row=linha - 1, rowOff=oy),
                               to=AnchorMarker(col=0, colOff=ox + MINI * PX, row=linha - 1, rowOff=oy + MINI * PX))
    ws.add_image(img)


def cabecalho(ws, colunas, cor=AZUL):
    for c, (t, w) in enumerate(colunas, 1):
        cel = ws.cell(row=1, column=c, value=t)
        cel.font, cel.fill, cel.alignment = cab_font, PatternFill("solid", fgColor=cor), centro
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"


def escrever(ws, r, valores, zebra=True):
    for c, v in enumerate(valores, 1):
        cel = ws.cell(row=r, column=c, value=v)
        cel.font, cel.alignment, cel.border = normal, quebra, borda
        if zebra and r % 2 == 0:
            cel.fill = PatternFill("solid", fgColor=CINZA)


def codigo_texto(cel):
    cel.number_format = "@"
    cel.font, cel.alignment = negrito, centro


wb = Workbook()

# ---------------------------------------------------------------------------
# PRODUTOS
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = "PRODUTOS"
COLS = [("IMAGEM", 7.5), ("CÓDIGO DO PRODUTO", 11), ("DESCRIÇÃO", 50), ("REFERÊNCIA DO FABRICANTE", 18),
        ("MARCA / FABRICANTE", 16), ("INFORMAÇÕES COMPLEMENTARES DO FORNECEDOR", 50), ("SALDO", 9),
        ("STATUS DA FOTO", 17), ("CATEGORIA", 22), ("ARQUIVO IMAGEM (SITE / CATÁLOGO)", 18),
        ("POST INSTAGRAM", 30), ("DESCRIÇÃO NO CATÁLOGO", 40), ("DIVERGÊNCIAS / CONFERIR", 60)]
cabecalho(ws, COLS)
ordem_status = {"COM FOTO": 0, "FOTO DE PRODUTO EQUIVALENTE": 1, "FOTO NÃO CONFERE": 2, "SEM FOTO": 3}
linhas_ord = sorted(linhas, key=lambda l: (ordem_status[l["status_foto"]], l["categoria"],
                                           ordem_natural(l["descricao"])))
for r, l in enumerate(linhas_ord, start=2):
    b = base.get(l["codigo"], {})
    divs = [f"[{t}] {a}" for t, a in sorted(l["alertas"], key=lambda x: PRIORIDADE.index(x[0]))]
    escrever(ws, r, [None, l["codigo"], l["descricao"], l["ref"], l["marca"], l["info"], l["saldo"],
                     l["status_foto"], l["categoria"], b.get("arquivo", ""), b.get("post_instagram", ""),
                     l["desc_catalogo"], "\n".join(divs)])
    codigo_texto(ws.cell(row=r, column=2))
    ws.cell(row=r, column=7).number_format = "#,##0.###"
    st = ws.cell(row=r, column=8)
    st.fill, st.alignment, st.font = PatternFill("solid", fgColor=COR_STATUS[l["status_foto"]]), centro, negrito
    if any(t not in INFORMATIVOS for t, _ in l["alertas"]):
        ws.cell(row=r, column=13).font = vermelho
    ws.row_dimensions[r].height = 27
    if b.get("img_mestre"):
        por_imagem(ws, b["img_mestre"], r, COLS[0][1])
ult_prod = len(linhas_ord) + 1
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{ult_prod}"

# ---------------------------------------------------------------------------
# SEM FOTO (BUSCAR) — ordenado por saldo: o que tem mais estoque vem primeiro
# ---------------------------------------------------------------------------
wsf = wb.create_sheet("SEM FOTO (BUSCAR)")
cabecalho(wsf, [("CÓDIGO", 10), ("DESCRIÇÃO", 55), ("REFERÊNCIA", 18), ("FABRICANTE", 18), ("SALDO", 9),
                ("CATEGORIA", 24), ("STATUS", 17), ("TERMO DE BUSCA SUGERIDO", 60)], cor=VERMELHO)
sem = sorted([l for l in linhas if l["status_foto"] in ("SEM FOTO", "FOTO NÃO CONFERE")],
             key=lambda l: -(l["saldo"] or 0))
for r, l in enumerate(sem, start=2):
    termo = " ".join(x for x in (l["marca"], l["ref"], re.sub(r"\s+", " ", l["descricao"])[:60]) if x)
    escrever(wsf, r, [l["codigo"], l["descricao"], l["ref"], l["marca"], l["saldo"], l["categoria"],
                      l["status_foto"], termo])
    codigo_texto(wsf.cell(row=r, column=1))
    wsf.cell(row=r, column=5).number_format = "#,##0.###"
wsf.auto_filter.ref = f"A1:H{len(sem) + 1}"

# ---------------------------------------------------------------------------
# DIVERGÊNCIAS
# ---------------------------------------------------------------------------
wd = wb.create_sheet("DIVERGÊNCIAS")
cabecalho(wd, [("CÓDIGO", 10), ("DESCRIÇÃO (PLANILHA)", 50), ("TIPO", 30), ("DETALHE / O QUE CONFERIR", 90),
               ("STATUS DA FOTO", 17)], cor=VERMELHO)
divs = sorted(((t, a, l) for l in linhas for t, a in l["alertas"]),
              key=lambda x: (PRIORIDADE.index(x[0]), x[2]["codigo"]))
for r, (t, a, l) in enumerate(divs, start=2):
    escrever(wd, r, [l["codigo"], l["descricao"], t + (" (informativo)" if t in INFORMATIVOS else ""), a,
                     l["status_foto"]])
    codigo_texto(wd.cell(row=r, column=1))
    if t not in INFORMATIVOS:
        wd.cell(row=r, column=3).font = Font(name=FONTE, size=9, bold=True, color="B00020")
wd.auto_filter.ref = f"A1:E{len(divs) + 1}"

# ---------------------------------------------------------------------------
# CATÁLOGO FORA DA LISTA
# ---------------------------------------------------------------------------
wc = wb.create_sheet("CATÁLOGO FORA DA LISTA")
COLC = [("IMAGEM", 7.5), ("CÓD. NO CATÁLOGO", 12), ("DESCRIÇÃO NO CATÁLOGO", 50), ("CATEGORIA", 24),
        ("MOTIVO", 42), ("POSSÍVEL CÓDIGO NA LISTA (pela descrição)", 70), ("PÁG. CATÁLOGO", 10)]
cabecalho(wc, COLC)
for r, f in enumerate(sorted(fora, key=lambda f: (f["categoria"], ordem_natural(f["descricao"]))), start=2):
    escrever(wc, r, [None, f["codigo"] or "S/ CÓDIGO", f["descricao"], f["categoria"], f["motivo"], f["sugestao"],
                     f["pagina"]])
    codigo_texto(wc.cell(row=r, column=2))
    wc.row_dimensions[r].height = 27
    if f.get("img_mestre"):
        por_imagem(wc, f["img_mestre"], r, COLC[0][1])
wc.auto_filter.ref = f"A1:G{len(fora) + 1}"

# ---------------------------------------------------------------------------
# IMAGENS ENVIADAS — candidatos na lista para o usuário escolher o código
# ---------------------------------------------------------------------------
we = wb.create_sheet("IMAGENS ENVIADAS")
COLE = [("IMAGEM", 7.5), ("IMAGEM ENVIADA (identificação)", 40), ("CÓDIGO CANDIDATO", 12),
        ("DESCRIÇÃO NA LISTA", 55), ("FABRICANTE", 16), ("SALDO", 9), ("É ESTE? (marque X)", 12)]
cabecalho(we, COLE)
r = 2
for e in enviadas:
    cands = e["candidatos"] or [{"codigo": "", "descricao": "Nenhum candidato encontrado na lista", "marca": "",
                                 "saldo": None}]
    for i, c in enumerate(cands):
        escrever(we, r, [None, e["descricao"] if i == 0 else "", c["codigo"], c["descricao"], c["marca"], c["saldo"],
                         ""], zebra=False)
        codigo_texto(we.cell(row=r, column=3))
        we.cell(row=r, column=7).fill = PatternFill("solid", fgColor="FFF2CC")
        we.row_dimensions[r].height = 27
        if i == 0 and e.get("img_mestre"):
            por_imagem(we, e["img_mestre"], r, COLE[0][1])
            we.cell(row=r, column=2).font = negrito
        r += 1
    for c in range(1, len(COLE) + 1):
        we.cell(row=r - 1, column=c).border = Border(bottom=Side(style="medium", color=AZUL))

# ---------------------------------------------------------------------------
# RESUMO (fórmulas)
# ---------------------------------------------------------------------------
wr = wb.create_sheet("RESUMO")
wr.column_dimensions["A"].width = 38
for col in "BCDEF":
    wr.column_dimensions[col].width = 16
titulos = ["CATEGORIA", "Nº DE ITENS", "COM FOTO", "SEM FOTO", "% COM FOTO", "SALDO SEM FOTO"]
for c, t in enumerate(titulos, 1):
    cel = wr.cell(row=1, column=c, value=t)
    cel.font, cel.fill, cel.alignment = cab_font, PatternFill("solid", fgColor=AZUL), centro
cats = sorted({l["categoria"] for l in linhas})
P = f"PRODUTOS!$I$2:$I${ult_prod}"
S_ = f"PRODUTOS!$H$2:$H${ult_prod}"
Q = f"PRODUTOS!$G$2:$G${ult_prod}"
for i, cat in enumerate(cats, start=2):
    wr.cell(row=i, column=1, value=cat).font = normal
    wr.cell(row=i, column=2, value=f"=COUNTIF({P},A{i})")
    wr.cell(row=i, column=3, value=f'=COUNTIFS({P},A{i},{S_},"COM FOTO")+COUNTIFS({P},A{i},{S_},"FOTO DE PRODUTO EQUIVALENTE")')
    wr.cell(row=i, column=4, value=f"=B{i}-C{i}")
    wr.cell(row=i, column=5, value=f"=IF(B{i}=0,0,C{i}/B{i})").number_format = "0.0%"
    wr.cell(row=i, column=6, value=f'=SUMIFS({Q},{P},A{i},{S_},"SEM FOTO")+SUMIFS({Q},{P},A{i},{S_},"FOTO NÃO CONFERE")')
    wr.cell(row=i, column=6).number_format = "#,##0"
    for c in range(2, 7):
        wr.cell(row=i, column=c).font = normal
t = len(cats) + 2
wr.cell(row=t, column=1, value="TOTAL").font = negrito
for c, col in zip(range(2, 7), "BCDEF"):
    f = f"=SUM({col}2:{col}{t - 1})" if col != "E" else f"=IF(B{t}=0,0,C{t}/B{t})"
    cel = wr.cell(row=t, column=c, value=f)
    cel.font = negrito
    cel.number_format = "0.0%" if col == "E" else "#,##0"
t += 2
extras = [("Itens com foto pelo código", f'=COUNTIF({S_},"COM FOTO")'),
          ("Itens com foto de produto equivalente (conferir)", f'=COUNTIF({S_},"FOTO DE PRODUTO EQUIVALENTE")'),
          ("Itens em que a foto do catálogo NÃO confere", f'=COUNTIF({S_},"FOTO NÃO CONFERE")'),
          ("Divergências (sem contar informativas)", f'=COUNTIF(DIVERGÊNCIAS!$C$2:$C${len(divs) + 1},"<>*informativo*")'),
          ("Produtos do catálogo fora da lista", f"=COUNTA('CATÁLOGO FORA DA LISTA'!$C$2:$C${len(fora) + 1})")]
for rot, v in extras:
    wr.cell(row=t, column=1, value=rot).font = normal
    wr.cell(row=t, column=2, value=v).font = negrito
    t += 1
wr.cell(row=t + 1, column=1, value=f"Fonte: {ORIGEM} (lista do sistema) cruzada com o catálogo PDF e as imagens "
                                   "enviadas. SALDO = quantidade em estoque informada na planilha.").font = Font(
    name=FONTE, size=8, italic=True, color="666666")

# ---------------------------------------------------------------------------
# LEIA-ME
# ---------------------------------------------------------------------------
wl = wb.create_sheet("LEIA-ME")
wl.column_dimensions["A"].width = 30
wl.column_dimensions["B"].width = 115
texto = [
    ("O QUE É", f"Sua lista de produtos ({ORIGEM}) com a foto tratada de cada item que existe no catálogo, "
                "além das divergências encontradas entre a lista e o catálogo."),
    ("IMAGEM", "Miniatura 30x30 px antes do código. Acompanha a linha ao filtrar/ordenar. A foto em alta está na "
               "pasta png_transparente_1200 com o nome = código."),
    ("STATUS DA FOTO", "COM FOTO: mesmo código do catálogo (ou item sem código no catálogo conferido pela "
                       "descrição/referência) · FOTO DE PRODUTO EQUIVALENTE: descrição igual a outro item do "
                       "catálogo — conferir · FOTO NÃO CONFERE: o catálogo usa o código, mas a marca é outra; "
                       "a foto não foi usada · SEM FOTO: buscar no site do fabricante (aba SEM FOTO)."),
    ("REFERÊNCIA / MARCA", "Valem as da sua planilha. Quando vazias, foram completadas com o catálogo "
                           "(marcado como informativo). Diferenças aparecem em DIVERGÊNCIAS."),
    ("INFORMAÇÕES COMPLEMENTARES", "Vêm do catálogo (medidas, aplicações, compatibilidades) para os itens com "
                                   "foto. Conferir a ficha do fornecedor antes de publicar especificações críticas."),
    ("SEM FOTO (BUSCAR)", "Itens sem foto ordenados pelo SALDO (mais estoque primeiro) com um termo de busca "
                          "pronto. Com a rede do ambiente liberada, as fotos podem ser buscadas nos sites dos "
                          "fabricantes."),
    ("DIVERGÊNCIAS", "Marca diferente, descrição diferente para o mesmo código, código repetido no catálogo, "
                     "referência diferente, descrições iguais com códigos diferentes (possível cadastro duplicado)."),
    ("CATÁLOGO FORA DA LISTA", "Produtos do catálogo que não estão na sua lista (a lista traz só itens com saldo). "
                               "Quando existe um item parecido na lista, o código provável aparece ao lado."),
    ("IMAGENS ENVIADAS", "Fotos enviadas sem código: marque com X o código correto para que a foto seja "
                         "vinculada na próxima geração."),
]
for i, (k, v) in enumerate(texto, start=1):
    a, b = wl.cell(row=i, column=1, value=k), wl.cell(row=i, column=2, value=v)
    a.font, b.font = Font(name=FONTE, size=10, bold=True, color=AZUL), Font(name=FONTE, size=10)
    a.alignment = b.alignment = Alignment(vertical="top", wrap_text=True)
    wl.row_dimensions[i].height = 44

XLSX.parent.mkdir(parents=True, exist_ok=True)
wb.save(XLSX)
print(f"{XLSX.name}: {len(linhas)} produtos | {sum(1 for l in linhas if base.get(l['codigo']))} com foto | "
      f"{len(sem)} sem foto | {len(divs)} divergências | {len(fora)} do catálogo fora da lista")
