"""Cruza a lista de produtos do sistema (xls/xlsx/csv) com as fotos tratadas do catálogo.

Regras:
  1. Foto pelo CÓDIGO (planilha x catálogo). Quando o catálogo usa o mesmo código para 2+ produtos,
     escolhe o de descrição mais parecida.
  2. Sem código no catálogo: foto de produto EQUIVALENTE só quando a descrição é praticamente igual
     (mesmos números/especificações) — marcada para conferência.
  3. Referência e marca: valem as da planilha; se vazias, usa as do catálogo (sinalizado);
     se diferentes, sinaliza.
  4. Divergências: especificação diferente para o mesmo código, códigos do catálogo fora da lista
     (com sugestão de código pela descrição), descrições idênticas com códigos diferentes.

Saída em <WORK_OUT>/:
  cruzamento.json      -> todas as linhas da planilha + status da foto + divergências
  base_produtos.json   -> só as linhas com foto (formato usado pelos scripts 03 e 04)
  fora_da_lista.json   -> produtos do catálogo que não estão na planilha

Uso: python 06_cruzar_planilha.py <planilha.xls|xlsx|csv> <WORK_CATALOGO> <WORK_OUT>
"""
import csv
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

from rapidfuzz import fuzz, process

PLANILHA, WORK_CAT, WORK_OUT = map(Path, sys.argv[1:4])
WORK_OUT.mkdir(parents=True, exist_ok=True)


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"\s+", " ", s).strip()


def compacto(s):
    return re.sub(r"[^A-Z0-9]", "", norm(s))


def numeros(s):
    """Conjunto de números da descrição (capacitância, tensão, corrente, medidas...). 1/3 é um número só."""
    out = set()
    texto = re.sub(r"\bREF[-:.]?\s*\S+", " ", norm(s))
    for n in re.findall(r"\d+(?:[.,]\d+)?(?:/\d+)?", texto):
        if "/" in n:
            out.add(n)
            continue
        n = n.replace(",", ".")
        if re.fullmatch(r"\d{1,3}\.\d{3}", n):  # 7.000 -> 7000
            n = n.replace(".", "")
        if "." in n:
            n = n.rstrip("0").rstrip(".")
        out.add(n.lstrip("0") or "0")
    return out


def vazio(v):
    return str(v or "").strip() in ("", "..", ".", "-", "0")


# ---------------------------------------------------------------------------
# Leitura da planilha (detecta as colunas pelo cabeçalho)
# ---------------------------------------------------------------------------
def ler_planilha(path):
    suf = path.suffix.lower()
    if suf == ".xls":
        import xlrd
        sh = xlrd.open_workbook(str(path)).sheet_by_index(0)
        linhas = [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
    elif suf in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True, read_only=True).worksheets[0]
        linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, encoding="utf-8-sig") as fh:
            linhas = list(csv.reader(fh))
    # cabeçalho = primeira linha que tenha "COD" e "DESC"
    for i, ln in enumerate(linhas[:20]):
        cab = [norm(c) for c in ln]
        if any(c.startswith("COD") for c in cab) and any("DESC" in c for c in cab):
            break
    else:
        raise SystemExit("Cabeçalho com CÓDIGO e DESCRIÇÃO não encontrado.")

    def col(*chaves):
        return next((j for j, c in enumerate(cab) if any(c.startswith(k) or k in c for k in chaves)), None)

    ic, idesc = col("COD"), col("DESC")
    iref, imarca, isaldo = col("REFER", "REF."), col("FABRIC", "MARCA"), col("SALDO", "ESTOQUE", "QTD")
    dados = []
    for ln in linhas[i + 1:]:
        if ic >= len(ln) or vazio(ln[ic]):
            continue
        cod = ln[ic]
        cod = str(int(cod)) if isinstance(cod, float) and cod.is_integer() else str(cod).strip()
        g = lambda j: (ln[j] if j is not None and j < len(ln) else "")
        dados.append({"codigo": cod, "descricao": str(g(idesc)).strip(),
                      "ref": "" if vazio(g(iref)) else str(g(iref)).strip(),
                      "marca": "" if vazio(g(imarca)) else str(g(imarca)).strip(),
                      "saldo": g(isaldo) if isinstance(g(isaldo), (int, float)) else None})
    print(f"Planilha: {len(dados)} produtos | colunas: código={cab[ic]}, descrição={cab[idesc]}, "
          f"ref={cab[iref] if iref is not None else '-'}, marca={cab[imarca] if imarca is not None else '-'}, "
          f"saldo={cab[isaldo] if isaldo is not None else '-'}")
    return dados


# ---------------------------------------------------------------------------
# Categoria sugerida para itens sem correspondência no catálogo
# ---------------------------------------------------------------------------
CATEGORIAS = [
    ("EPI / Segurança", r"\b(LUVA|OCULOS|BOTINA|SAPATO|BOTA|TENIS|CAPACETE|PROTETOR AURICULAR|MASCARA|AVENTAL|"
                        r"CINTO|PERNEIRA|CAPA DE CHUVA)\b"),
    ("Capacitores", r"\bCAPACITOR"),
    ("Gases Refrigerantes", r"\b(GAS|FLUIDO REFRIG)\b|\bR-?(22|32|134A?|404A?|407C?|410A?|600A?|290|141B)\b"),
    ("Químicos / Solda e Limpeza", r"\b(SOLDA|FLUXO|VARETA|ALUMAX|TAPA FUGAS|SELANTE|DESUMIFICANTE|CAPIDRYL|"
                                   r"LIMPA|ANTICHAMA|COLA|FITA VEDA|SILICONE|OLEO|DESENGRAX|MACARICO|TINTA|WD-?40|"
                                   r"ALCOOL|DETERGENTE|MULTIUSO|ADITIVO|REFRICLEAN|DESCONTAMINANTE|GRAXA|"
                                   r"POLIURETANO|ESPUMA EXPANS|THINNER|VERNIZ|LUBRIFICANTE)"),
    ("Iluminação", r"\b(LUMINARIA|SPOT|PLAFON|FITA (DE )?LED|LED\b|LAMPADA|ARANDELA|REFLETOR|BOCAL|PAINEL LED|"
                   r"TRILHO ELETRIFICADO|PERFIL (DE )?LED)"),
    ("Ferramentas / Instrumentos", r"\b(CHAVE|ALICATE|MANIFOLD|FLANGEADOR|BOMBA DE VACUO|BOMBA VACUO|MULTIMETRO|"
                                   r"AMPERIMETRO|CAPACIMETRO|TERMOMETRO|VACUOMETRO|CURVADOR|CORTADOR|MOLA|PENTE|"
                                   r"MANGUEIRA DE CARGA|ESCAREADOR|SERRA|BROCA|MARTELO|TRENA|NIVEL|FERRO DE SOLDA|"
                                   r"MALA|MOCHILA|MANOMETRO|DETECTOR|ESTILETE|LIMA|TESOURA|SOQUETE|EXPANSOR|"
                                   r"PARAFUSADEIRA|FURADEIRA|MARTELETE|ESMERILHADEIRA|LIXADEIRA|SOPRADOR|"
                                   r"ASPIRADOR|DISCO|ESCOVA DE CARVAO|ESCOVA|JOGO|ESCADA|SACA|GARRA|PONTEIRA|"
                                   r"TALHADEIRA|FORMAO|ARCO)"),
    ("Hidráulica", r"\b(SIFAO|REDUCAO|COTOVELO|TEE|JOELHO|ESGUICHO|MANGOTE|ENGATE|VALVULA DE ESCOAMENTO|"
                   r"TUBO SOLDAVEL|CURVA 90|CURVA 45|LUVA SOLD|ADAPTADOR SOLD|CAIXA D.?AGUA|BOIA|CAP SOLD)"),
    ("Fixação e Acessórios", r"\b(ABR|ABRACADEIRA|ARRUELA|PARAFUSO|PF|BUCHA|PORCA|REBITE|TIRANTE|PREGO|"
                             r"CANTONEIRA|CHUMBADOR|BARRA ROSCADA|FITA HELLERMAN|CINTA|DOBRADICA|TRINCO|"
                             r"FECHADURA|CORDA|CADEADO)\b"),
    ("Eletrônicos / Informática", r"\b(CAMERA|MOUSE|TECLADO|FONE|PENDRIVE|CARREGADOR|ANTENA|BATERIA|PILHA|"
                                  r"FONTE|CABO DE DADOS|CABO USB|CABO HDMI|ADAPTADOR USB|CAIXA DE SOM|"
                                  r"CONTROLE UNIVERSAL TV|ROTEADOR|HD EXTERNO|CARTAO DE MEMORIA|DRIVE)\b"),
    ("Elétrica / Comandos", r"\b(DISJUNTOR|CONTATOR|CONTATORA|RELE TERMICO|TOMADA|PLUG|PLUGUE|INTERRUPTOR|CABO|FIO|"
                            r"DPS|QUADRO|SINALEIRO|BOTAO|TERMINAL|CONECTOR|FILTRO DE LINHA|EXTENSAO|FUSIVEL|"
                            r"REATOR|CANALETA|ELETRODUTO|CAIXA|CX|BORNE|CONDULETE|DIMMER|SINALIZADOR|SIRENE|"
                            r"MODULO|TRILHO|INDICADOR|ENTRADA|BARRAMENTO|ESPELHO|PLACA CEGA|CAMPAINHA|"
                            r"MINUTERIA|FOTOCELULA|RELE FOTO|SENSOR DE PRESENCA|ESTABILIZADOR|NOBREAK)"),
    ("Automação / Temporizadores", r"\b(TEMPORIZADOR|TIMER|RELE DE TEMPO|CONTROLADOR|PROGRAMADOR|PRESSOSTATO|"
                                   r"FALTA DE FASE)"),
    ("Filtros e Purificadores de Água", r"\b(REFIL|PURIFICADOR|FILTRO DE AGUA|FILTRO AGUA|ELEMENTO FILTRANTE)"),
    ("Peças para Lavadora e Bebedouro", r"\b(LAVADORA|LAVAD|BEBEDOURO|TORNEIRA|CESTO|AGITADOR|VARA DE SUSP|POLIA|"
                                        r"TANQUINHO|MICROONDAS|MICRO-ONDAS|FOGAO|ROLAMENTO|RETENTOR|TRANSMISS|"
                                        r"SUSPENSAO|CORREIA|EIXO|MANCAL|CARCACA|PISTAO|CONJ)"),
    ("Instalação de Ar-Condicionado", r"\b(TUBO|SUPORTE|ISOLAMENTO|ISOLANTE|ESPONJOSO|TUBEX|ELASTOMER|DRENO|"
                                      r"MANGUEIRA|CANO|CURVA|CONEXAO|UNIAO|FLANGE|ESPUMA|FITA PVC|MANTA|"
                                      r"FITA ALUMINIZADA|FITA)"),
    ("Refrigeração / Componentes", r"\b(COMPRESSOR|RELE|PROTETOR|TERMOSTATO|VALVULA|FILTRO SECADOR|FILTRO|MOTOR|"
                                   r"VENTILADOR|HELICE|SENSOR|PLACA|RESISTENCIA|EVAPORADOR|CONDENSADOR|"
                                   r"SERPENTINA|BORRACHA|GAXETA|CAPILAR|DAMPER|BANDEJA|RECIPIENTE|KIT|TRAFO|"
                                   r"TRANSFORMADOR|CONTROLE|ACUMULADOR|REGISTRO|VISOR|EXAUSTOR|COLETOR|BOMBA|"
                                   r"ELEMENTO|BLOCO|DOBRADICA|PUXADOR|TRAVA)"),
]


def categoria_sugerida(desc):
    d = norm(desc)
    for cat, pad in CATEGORIAS:
        if re.search(pad, d):
            return cat
    return "Outros"


ALIAS_MARCA = {"BW": "BRASWELD"}


def marca_equivalente(plan, cat, desc_plan=""):
    """Mesma marca? Considera apelidos (BW = BRASWELD) e marca do catálogo citada na descrição da planilha."""
    a, b = compacto(ALIAS_MARCA.get(norm(plan), plan)), compacto(ALIAS_MARCA.get(norm(cat), cat))
    if not a or not b or a in b or b in a or fuzz.ratio(a, b) >= 85:
        return True
    return compacto(cat) in compacto(desc_plan)


def ref_equivalente(plan, cat, desc_plan=""):
    a, b = compacto(plan), compacto(cat)
    return not a or not b or a in b or b in a or b in compacto(desc_plan)


VARIANTES = {"PRETO", "PRETA", "BRANCO", "BRANCA", "AZUL", "VERMELHO", "VERMELHA", "VERDE", "AMARELO", "AMARELA",
             "CINZA", "BLACK", "INOX", "FUME", "INCOLOR", "BEGE", "MARROM", "CAFE", "LARANJA", "ROSA", "LILAS"}

# Produtos do catálogo sem código, identificados na planilha pela descrição/referência (conferido manualmente)
MANUAL = {"053699": "motor_black", "054357": "luva_ve801", "054358": "luva_ve801", "054338": "oculos_summer"}

# Imagens enviadas sem código: palavras-chave para listar candidatos na planilha (o usuário escolhe)
BUSCA_ENVIADAS = {
    "user:1.jpg": [r"BOTINA", r"ELASTIC|ELESTIC"],
    "user:2.jpg": [r"BOTINA", r"AMARRAR|AMA |CADAR"],
    "user:3.jpg": [r"PROTETOR|DPS", r"RAIO|SURTO|TOMADA"],
    "user:4.jpg": [r"SENSOR", r"A/C|AR COND|DEGELO|DEG "],
    "user:5.jpg": [r"SENSOR", r"EVAP|SERPENT|DEGELO|10K"],
    "psd:contator_01_polo": [r"CONTATOR|CONTATORA", r"BIF|1P|1 POLO|MONO|25A"],
    "psd:KTM_1000_copiar": [r"KTM"],
}


# ---------------------------------------------------------------------------
def main():
    planilha = ler_planilha(PLANILHA)
    catalogo = json.load(open(WORK_CAT / "base_produtos.json"))
    cat_cod = defaultdict(list)
    for p in catalogo:
        if p["codigo"]:
            cat_cod[int(p["codigo"])].append(p)
    usados = set()  # id() dos produtos do catálogo usados pela planilha

    # índice para busca por descrição
    cat_com_foto = [p for p in catalogo if p.get("img")]
    cat_desc = [norm(p["descricao"]) for p in cat_com_foto]

    linhas = []
    for item in planilha:
        alertas, cat_p, status, metodo = [], None, "SEM FOTO", ""
        codigo_int = int(item["codigo"]) if item["codigo"].isdigit() else None
        cands = cat_cod.get(codigo_int, [])
        if cands:
            cat_p = max(cands, key=lambda c: (numeros(c["descricao"]) <= numeros(item["descricao"]),
                                              fuzz.token_set_ratio(norm(c["descricao"]), norm(item["descricao"]))))
            status, metodo = "COM FOTO", "Código igual ao do catálogo"
            if len(cands) > 1:
                outros = "; ".join(c["descricao"] for c in cands if c is not cat_p)
                alertas.append(("Código repetido no catálogo",
                                f"No catálogo este código também aparece em: {outros}. Foto escolhida pela descrição."))
            # especificação: números do catálogo que não aparecem na planilha
            faltam = numeros(cat_p["descricao"]) - numeros(item["descricao"])
            if faltam:
                alertas.append(("Descrição diferente para o mesmo código",
                                f"Catálogo: '{cat_p['descricao']}' | Planilha: '{item['descricao']}' "
                                f"(números só no catálogo: {', '.join(sorted(faltam))})"))
        else:
            # produto equivalente pela descrição (mesmos números, texto quase igual)
            achado = process.extractOne(norm(item["descricao"]), cat_desc, scorer=fuzz.token_sort_ratio,
                                        score_cutoff=90)
            manual = next((c for c in catalogo if c["familia"] == MANUAL.get(item["codigo"])), None)
            if manual:
                cat_p, status = manual, "COM FOTO"
                metodo = "Produto sem código no catálogo, identificado pela descrição/referência (conferido)"
            elif achado:
                cand = cat_com_foto[achado[2]]
                extras = set(norm(item["descricao"]).split()) - set(norm(cand["descricao"]).split())
                if (numeros(cand["descricao"]) == numeros(item["descricao"]) and not (extras & VARIANTES)
                        and marca_equivalente(item["marca"], cand.get("marca", ""), item["descricao"])):
                    cat_p, status = cand, "FOTO DE PRODUTO EQUIVALENTE"
                    origem = f"cód. {cand['codigo']}" if cand["codigo"] else "item sem código"
                    metodo = f"Descrição igual à do {origem} do catálogo (conferir)"
                    alertas.append(("Foto de produto equivalente",
                                    f"Foto do {origem} do catálogo '{cand['descricao']}' — conferir se é o mesmo item."))
        if cat_p is not None:
            usados.add(id(cat_p))

        # referência e marca: planilha manda; catálogo completa
        ref, marca = item["ref"], item["marca"]
        if cat_p is not None:
            if not ref and cat_p.get("ref"):
                ref = cat_p["ref"]
                alertas.append(("Referência vinda do catálogo", f"Referência vazia na planilha; usada a do catálogo: {ref}"))
            elif ref and cat_p.get("ref") and not ref_equivalente(ref, cat_p["ref"], item["descricao"]):
                alertas.append(("Referência diferente", f"Planilha: {ref} | Catálogo: {cat_p['ref']}"))
            if not marca and cat_p.get("marca"):
                marca = cat_p["marca"]
                alertas.append(("Marca vinda do catálogo", f"Fabricante vazio na planilha; usada a marca do catálogo: {marca}"))
            elif marca and cat_p.get("marca") and not marca_equivalente(marca, cat_p["marca"], item["descricao"]):
                alertas.append(("Marca diferente", f"Planilha: {marca} | Catálogo: {cat_p['marca']} — a foto do "
                                                   f"catálogo é de outra marca e NÃO foi usada."))
                status, metodo = "FOTO NÃO CONFERE", f"Foto do catálogo (cód. {cat_p['codigo']}) é da marca {cat_p['marca']}"

        linhas.append({
            **item, "ref": ref, "marca": marca, "status_foto": status, "metodo_foto": metodo,
            "cod_catalogo": cat_p["codigo"] if cat_p else "", "desc_catalogo": cat_p["descricao"] if cat_p else "",
            "img": cat_p["img"] if cat_p and status != "FOTO NÃO CONFERE" else None,
            "img_mestre": cat_p.get("img_mestre") if cat_p and status != "FOTO NÃO CONFERE" else None,
            "familia": cat_p["familia"] if cat_p else "", "familia_titulo": cat_p["familia_titulo"] if cat_p else "",
            "categoria": ({"Químicos / Solda e Vedação": "Químicos / Solda e Limpeza"}.get(cat_p["categoria"],
                                                                                        cat_p["categoria"])
                          if cat_p else categoria_sugerida(item["descricao"])),
            "info": cat_p["info"] if cat_p else "", "pagina": cat_p.get("pagina") if cat_p else None,
            "alertas": alertas,
        })

    # descrições idênticas com códigos diferentes (possível cadastro duplicado no sistema)
    por_desc = defaultdict(list)
    for ln in linhas:
        por_desc[compacto(ln["descricao"])].append(ln)
    for grupo in por_desc.values():
        if len({g["codigo"] for g in grupo}) > 1:
            for g in grupo:
                outros = ", ".join(o["codigo"] for o in grupo if o is not g)
                g["alertas"].append(("Mesma descrição com outro código",
                                     f"Descrição idêntica ao(s) código(s) {outros}: possível cadastro duplicado."))

    # imagens enviadas sem código: sugere candidatos na planilha
    desc_plan = [norm(l["descricao"]) for l in linhas]
    sugestoes = []
    for p in catalogo:
        if p["origem"] == "Imagem enviada":
            if p["codigo"] and any(l["codigo"].lstrip("0") == p["codigo"] for l in linhas):
                continue
            pads = BUSCA_ENVIADAS.get(p["img"], [])
            cand = [l for l in linhas if pads and all(re.search(pd, norm(l["descricao"])) for pd in pads)]
            sugestoes.append({"imagem": p["img"], "descricao": p["descricao"], "img_mestre": p.get("img_mestre"),
                              "candidatos": [{"codigo": l["codigo"], "descricao": l["descricao"], "marca": l["marca"],
                                              "saldo": l["saldo"]} for l in cand[:15]]})

    # produtos do catálogo fora da lista (com sugestão de código pela descrição)
    fora = []
    for p in catalogo:
        if p["origem"] != "Catálogo PDF" or id(p) in usados or any(
                l["cod_catalogo"] == p["codigo"] and l["desc_catalogo"] == p["descricao"] for l in linhas):
            continue
        motivo = "Código não está na lista (sem saldo na loja ou código diferente)"
        if p["codigo"] and int(p["codigo"]) in {int(l["codigo"]) for l in linhas if l["codigo"].isdigit()}:
            motivo = "Código existe na lista, mas para outro produto (código repetido no catálogo)"
        elif not p["codigo"]:
            motivo = "Produto sem código no catálogo e não identificado na lista"
        melhor = process.extractOne(norm(p["descricao"]), desc_plan, scorer=fuzz.token_sort_ratio)
        sug = ""
        if melhor and melhor[1] >= 85 and numeros(p["descricao"]) <= numeros(linhas[melhor[2]]["descricao"]):
            alvo = linhas[melhor[2]]
            sug = f"{alvo['codigo']} — {alvo['descricao']}"
            if alvo["marca"] and p.get("marca") and not marca_equivalente(alvo["marca"], p["marca"], alvo["descricao"]):
                sug += f" (marca {alvo['marca']} ≠ {p['marca']} do catálogo: produto similar, não o mesmo)"
        fora.append({"codigo": p["codigo"], "descricao": p["descricao"], "categoria": p["categoria"], "motivo": motivo,
                     "img_mestre": p.get("img_mestre"), "pagina": p.get("pagina"), "sugestao": sug})

    # base para os scripts 03/04 (somente com foto)
    base = []
    for ln in linhas:
        if ln["img"]:
            base.append({"codigo": ln["codigo"], "chave": ln["codigo"], "descricao": ln["descricao"],
                         "familia": ln["familia"], "familia_titulo": ln["familia_titulo"],
                         "categoria": ln["categoria"], "marca": ln["marca"], "ref": ln["ref"], "info": ln["info"],
                         "img": ln["img"], "pagina": ln["pagina"], "alertas": [a[1] for a in ln["alertas"]],
                         "origem": "Planilha do sistema"})

    json.dump(linhas, open(WORK_OUT / "cruzamento.json", "w"), ensure_ascii=False, indent=1)
    json.dump(base, open(WORK_OUT / "base_produtos.json", "w"), ensure_ascii=False, indent=1)
    json.dump(fora, open(WORK_OUT / "fora_da_lista.json", "w"), ensure_ascii=False, indent=1)
    json.dump(sugestoes, open(WORK_OUT / "sugestoes_imagens_enviadas.json", "w"), ensure_ascii=False, indent=1)

    from collections import Counter
    print("Status das fotos:", dict(Counter(l["status_foto"] for l in linhas)))
    print("Divergências:", dict(Counter(a[0] for l in linhas for a in l["alertas"])))
    print(f"Catálogo fora da lista: {len(fora)} (com sugestão de código: {sum(1 for f in fora if f['sugestao'])})")


if __name__ == "__main__":
    main()
