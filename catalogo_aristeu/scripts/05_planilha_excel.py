"""Gera a planilha de produtos com a coluna IMAGEM (miniatura 30x30 px) antes do código.

Uso: python 05_planilha_excel.py <WORK_DIR> <SAIDA_IMAGENS> <ARQUIVO_XLSX>
"""
import io
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image

WORK, SAIDA, XLSX = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
base = json.load(open(WORK / "base_produtos.json"))

FONTE = "Arial"
AZUL, VERMELHO, CINZA_CLARO = "10284F", "E81820", "F3F5F9"
PX_EMU = 9525
MINI = 30  # tamanho exibido da miniatura (px), conforme pedido

cab = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
fill_cab = PatternFill("solid", fgColor=AZUL)
normal = Font(name=FONTE, size=9)
negrito = Font(name=FONTE, size=9, bold=True)
alerta_font = Font(name=FONTE, size=9, color="B00020")
borda = Border(bottom=Side(style="thin", color="D9DEE7"))
quebra = Alignment(vertical="center", wrap_text=True)
centro = Alignment(horizontal="center", vertical="center")

wb = Workbook()

# ---------------------------------------------------------------------------
# Aba PRODUTOS
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = "PRODUTOS"
colunas = [
    ("IMAGEM", 5.6),
    ("CÓDIGO DO PRODUTO", 11),
    ("DESCRIÇÃO", 52),
    ("REFERÊNCIA DO FABRICANTE", 18),
    ("MARCA", 14),
    ("INFORMAÇÕES COMPLEMENTARES DO FORNECEDOR", 60),
    ("CATEGORIA", 24),
    ("ARQUIVO IMAGEM (SITE / CATÁLOGO)", 30),
    ("POST INSTAGRAM", 34),
    ("TRATAMENTO DA IMAGEM", 30),
    ("ALERTAS / CONFERIR", 60),
    ("ORIGEM", 16),
]
for c, (titulo, larg) in enumerate(colunas, 1):
    cel = ws.cell(row=1, column=c, value=titulo)
    cel.font, cel.fill = cab, fill_cab
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(c)].width = larg
ws.row_dimensions[1].height = 32

mini_cache = {}


def miniatura(mestre):
    """PNG 120x120 (4x a exibição) para a miniatura ficar nítida com zoom no Excel."""
    if mestre not in mini_cache:
        im = Image.open(SAIDA / "_mestres" / mestre).convert("RGBA")
        im.thumbnail((112, 112), Image.LANCZOS)
        tela = Image.new("RGBA", (120, 120), (255, 255, 255, 0))
        tela.alpha_composite(im, ((120 - im.width) // 2, (120 - im.height) // 2))
        buf = io.BytesIO()
        tela.save(buf, format="PNG", optimize=True)
        mini_cache[mestre] = buf.getvalue()
    return io.BytesIO(mini_cache[mestre])


for r, p in enumerate(base, start=2):
    origem = f"Catálogo pág. {p['pagina']}" if p["pagina"] else p["origem"]
    valores = [None, p["codigo"] or "S/ CÓDIGO", p["descricao"], p["ref"], p["marca"], p["info"], p["categoria"],
               p["arquivo"], p.get("post_instagram", ""), p["img_metodo"], " | ".join(p["alertas"]), origem]
    for c, v in enumerate(valores, 1):
        cel = ws.cell(row=r, column=c, value=v)
        cel.font, cel.alignment, cel.border = normal, quebra, borda
    ws.cell(row=r, column=2).font = negrito
    ws.cell(row=r, column=2).alignment = centro
    if p["codigo"].isdigit():
        ws.cell(row=r, column=2).value = int(p["codigo"])
        ws.cell(row=r, column=2).number_format = "0"
    if p["alertas"]:
        ws.cell(row=r, column=11).font = alerta_font
    if r % 2 == 0:
        for c in range(1, len(colunas) + 1):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=CINZA_CLARO)
    ws.row_dimensions[r].height = 27  # ~36 px: miniatura 30 px + respiro

    img = XLImage(miniatura(p["img_mestre"]))
    img.width = img.height = MINI
    col_px = int(colunas[0][1] * 7 + 5)
    off_x = max(0, (col_px - MINI) // 2) * PX_EMU
    off_y = 3 * PX_EMU
    img.anchor = OneCellAnchor(_from=AnchorMarker(col=0, colOff=off_x, row=r - 1, rowOff=off_y),
                               ext=XDRPositiveSize2D(MINI * PX_EMU, MINI * PX_EMU))
    ws.add_image(img)

ultima = len(base) + 1
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{ultima}"
ws.sheet_view.zoomScale = 110

# ---------------------------------------------------------------------------
# Aba ALERTAS (conferência de cadastro antes de publicar / comprar)
# ---------------------------------------------------------------------------
wa = wb.create_sheet("ALERTAS CADASTRO")
cab_a = ["CÓDIGO", "DESCRIÇÃO", "TIPO", "ALERTA / O QUE CONFERIR"]
for c, t in enumerate(cab_a, 1):
    cel = wa.cell(row=1, column=c, value=t)
    cel.font, cel.fill, cel.alignment = cab, PatternFill("solid", fgColor=VERMELHO), centro
for col, w in zip("ABCD", (11, 55, 26, 90)):
    wa.column_dimensions[col].width = w


def tipo(a):
    if a.startswith("CÓDIGO REPETIDO"):
        return "Código repetido"
    if a.startswith("Descrição corrigida"):
        return "Descrição corrigida"
    if "sem código" in a.lower() or a.startswith("Sem código"):
        return "Sem código"
    if "duplicado" in a:
        return "Possível cadastro duplicado"
    return "Conferir informação"


r = 2
prioridade = {"Código repetido": 0, "Possível cadastro duplicado": 1, "Conferir informação": 2, "Sem código": 3,
              "Descrição corrigida": 4}
linhas = sorted(((tipo(a), p, a) for p in base for a in p["alertas"]), key=lambda t: (prioridade[t[0]], t[1]["codigo"]))
for t, p, a in linhas:
    for c, v in enumerate([int(p["codigo"]) if p["codigo"].isdigit() else "S/ CÓDIGO", p["descricao"], t, a], 1):
        cel = wa.cell(row=r, column=c, value=v)
        cel.font, cel.alignment, cel.border = normal, quebra, borda
    r += 1
wa.freeze_panes = "A2"
wa.auto_filter.ref = f"A1:D{r - 1}"

# ---------------------------------------------------------------------------
# Aba RESUMO (fórmulas — atualiza se a aba PRODUTOS for editada)
# ---------------------------------------------------------------------------
wr = wb.create_sheet("RESUMO")
wr.column_dimensions["A"].width = 38
for col in "BCD":
    wr.column_dimensions[col].width = 18
titulos = ["CATEGORIA", "Nº DE ITENS", "ITENS COM ALERTA", "% COM ALERTA"]
for c, t in enumerate(titulos, 1):
    cel = wr.cell(row=1, column=c, value=t)
    cel.font, cel.fill, cel.alignment = cab, fill_cab, centro
cats = sorted({p["categoria"] for p in base})
for i, cat in enumerate(cats, start=2):
    wr.cell(row=i, column=1, value=cat).font = normal
    wr.cell(row=i, column=2, value=f"=COUNTIF(PRODUTOS!$G$2:$G${ultima},A{i})").font = normal
    wr.cell(row=i, column=3, value=f'=COUNTIFS(PRODUTOS!$G$2:$G${ultima},A{i},PRODUTOS!$K$2:$K${ultima},"?*")').font = normal
    cel = wr.cell(row=i, column=4, value=f"=IF(B{i}=0,0,C{i}/B{i})")
    cel.font, cel.number_format = normal, "0.0%"
t = len(cats) + 2
wr.cell(row=t, column=1, value="TOTAL").font = negrito
wr.cell(row=t, column=2, value=f"=SUM(B2:B{t - 1})").font = negrito
wr.cell(row=t, column=3, value=f"=SUM(C2:C{t - 1})").font = negrito
cel = wr.cell(row=t, column=4, value=f"=IF(B{t}=0,0,C{t}/B{t})")
cel.font, cel.number_format = negrito, "0.0%"
t += 2
extras = [("Itens sem código", f'=COUNTIF(PRODUTOS!$B$2:$B${ultima},"S/ CÓDIGO")'),
          ("Imagens de produto distintas", len({p['img'] for p in base})),
          ("Posts de Instagram gerados", len({p.get('post_instagram') for p in base}))]
for rot, v in extras:
    wr.cell(row=t, column=1, value=rot).font = normal
    wr.cell(row=t, column=2, value=v).font = normal
    t += 1
wr.cell(row=t, column=1, value="Fonte: catálogo 'oficial_livro_paginas_impares.pdf' + imagens enviadas "
                                "(fotos e PSD). Contagens de imagens/posts calculadas na geração do arquivo.").font = Font(
    name=FONTE, size=8, italic=True, color="666666")

# ---------------------------------------------------------------------------
# Aba LEIA-ME
# ---------------------------------------------------------------------------
wl = wb.create_sheet("LEIA-ME")
wl.column_dimensions["A"].width = 32
wl.column_dimensions["B"].width = 110
texto = [
    ("O QUE É", "Base de produtos extraída do catálogo Aristeu Refrigeração (PDF) com foto tratada de cada item."),
    ("IMAGEM", "Miniatura 30x30 px (clique na imagem para ver maior; ela foi inserida com resolução 4x para ficar "
               "nítida com zoom)."),
    ("ARQUIVO IMAGEM", "Nome do PNG 1200x1200 com fundo transparente na pasta imagens/png_transparente_1200 — "
                       "pronto para o site e para montar catálogo. O nome do arquivo é o código do produto."),
    ("POST INSTAGRAM", "Arte 1080x1350 (formato retrato 4:5) na pasta imagens/instagram_1080x1350 — um post por "
                       "produto/família, com logo, códigos e WhatsApp das lojas."),
    ("TRATAMENTO", "Fotos do catálogo ampliadas 4x por IA (Real-ESRGAN). Onde a foto não tinha recorte, o fundo foi "
                   "removido por IA (ISNet). As fotos originais do catálogo são pequenas (≈100–400 px); para "
                   "zoom de e-commerce acima de 1200 px, o ideal é pedir as fotos originais ao fornecedor."),
    ("INFORMAÇÕES COMPLEMENTARES", "Dados do próprio catálogo (medidas, aplicações, compatibilidades, referências) + "
                                   "descrição técnica geral do tipo de produto. Conferir ficha técnica do fornecedor "
                                   "antes de publicar especificações críticas."),
    ("ALERTAS / CONFERIR", "Problemas encontrados no catálogo: mesmo código usado para produtos diferentes, "
                           "descrições trocadas, produtos sem código e correções de digitação. Ver aba "
                           "ALERTAS CADASTRO — corrigir no sistema evita venda e compra do item errado."),
    ("ITENS 'S/ CÓDIGO'", "Aparecem no catálogo ou nas imagens enviadas sem código. Preencher a coluna CÓDIGO com "
                          "o código do sistema e renomear o arquivo PNG correspondente."),
    ("COMO EDITAR", "Pode editar/preencher as colunas B a G livremente; a aba RESUMO recalcula sozinha."),
]
for i, (k, v) in enumerate(texto, start=1):
    a = wl.cell(row=i, column=1, value=k)
    b = wl.cell(row=i, column=2, value=v)
    a.font, b.font = Font(name=FONTE, size=10, bold=True, color=AZUL), Font(name=FONTE, size=10)
    a.alignment = b.alignment = Alignment(vertical="top", wrap_text=True)
    wl.row_dimensions[i].height = 42

XLSX.parent.mkdir(parents=True, exist_ok=True)
wb.save(XLSX)
print(f"{XLSX}: {len(base)} produtos, {r - 2} alertas")
